from flask import Flask, request, jsonify, send_file
import fitz
import json

from io import BytesIO

from docx import Document
from docx.shared import Pt
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH


app = Flask(__name__)


# ============================================================
# CONFIG
# ============================================================

CHUNK_SIZE = 40

DEFAULT_FONT_NAME = "Arial"
DEFAULT_FONT_SIZE = 10

MIN_VERTICAL_GAP_PT = 0
MAX_VERTICAL_GAP_PT = 40


# ============================================================
# ROOT
# ============================================================

@app.route("/", methods=["GET"])
def home():
    return {
        "status": "ok",
        "service": "PDF Parser"
    }


# ============================================================
# HELPERS
# ============================================================

def clamp(value, minimum, maximum):
    return max(
        minimum,
        min(
            maximum,
            value
        )
    )


def extract_pdf_elements(document):
    pages = []
    elements = []

    element_counter = 1

    for page_index in range(
        len(document)
    ):
        page = document[
            page_index
        ]

        blocks = page.get_text(
            "blocks"
        )

        page_elements = []

        for block in blocks:

            if len(block) < 5:
                continue

            x0 = float(
                block[0]
            )

            y0 = float(
                block[1]
            )

            x1 = float(
                block[2]
            )

            y1 = float(
                block[3]
            )

            text = str(
                block[4]
            ).strip()

            if not text:
                continue

            element = {
                "id":
                    f"pdf_{element_counter}",

                "page":
                    page_index + 1,

                "type":
                    "text_block",

                "text":
                    text,

                "bbox": {
                    "x0":
                        round(
                            x0,
                            2
                        ),

                    "y0":
                        round(
                            y0,
                            2
                        ),

                    "x1":
                        round(
                            x1,
                            2
                        ),

                    "y1":
                        round(
                            y1,
                            2
                        )
                }
            }

            elements.append(
                element
            )

            page_elements.append(
                element
            )

            element_counter += 1

        pages.append({
            "page":
                page_index + 1,

            "width":
                round(
                    page.rect.width,
                    2
                ),

            "height":
                round(
                    page.rect.height,
                    2
                ),

            "element_count":
                len(
                    page_elements
                ),

            "elements":
                page_elements
        })

    return pages, elements


def parse_translations(raw_value):
    if raw_value is None:
        return {}

    if isinstance(
        raw_value,
        bytes
    ):
        raw_value = raw_value.decode(
            "utf-8"
        )

    if isinstance(
        raw_value,
        str
    ):
        data = json.loads(
            raw_value
        )
    else:
        data = raw_value

    translations = data.get(
        "translations",
        []
    )

    result = {}

    for item in translations:

        element_id = item.get(
            "id"
        )

        text = item.get(
            "text"
        )

        if not element_id:
            continue

        if text is None:
            text = ""

        result[
            str(element_id)
        ] = str(text)

    return result


def configure_section(
    section,
    page_width,
    page_height
):
    section.page_width = Pt(
        page_width
    )

    section.page_height = Pt(
        page_height
    )

    # bbox değerlerini doğrudan
    # kullanabilmek için kenarları
    # mümkün olduğunca sıfıra yaklaştırıyoruz.
    section.top_margin = Pt(
        0
    )

    section.bottom_margin = Pt(
        0
    )

    section.left_margin = Pt(
        0
    )

    section.right_margin = Pt(
        0
    )

    section.header_distance = Pt(
        0
    )

    section.footer_distance = Pt(
        0
    )


def add_text_block(
    document,
    element,
    translated_text,
    page_width,
    previous_bottom
):
    bbox = element[
        "bbox"
    ]

    x0 = float(
        bbox["x0"]
    )

    y0 = float(
        bbox["y0"]
    )

    x1 = float(
        bbox["x1"]
    )

    y1 = float(
        bbox["y1"]
    )

    paragraph = document.add_paragraph()

    paragraph.paragraph_format.left_indent = Pt(
        max(
            0,
            x0
        )
    )

    right_indent = max(
        0,
        page_width - x1
    )

    paragraph.paragraph_format.right_indent = Pt(
        right_indent
    )

    if previous_bottom is None:
        vertical_gap = y0
    else:
        vertical_gap = (
            y0 -
            previous_bottom
        )

    vertical_gap = clamp(
        vertical_gap,
        MIN_VERTICAL_GAP_PT,
        MAX_VERTICAL_GAP_PT
    )

    paragraph.paragraph_format.space_before = Pt(
        vertical_gap
    )

    paragraph.paragraph_format.space_after = Pt(
        0
    )

    paragraph.paragraph_format.line_spacing = 1

    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.LEFT
    )

    run = paragraph.add_run(
        translated_text
    )

    run.font.name = (
        DEFAULT_FONT_NAME
    )

    run.font.size = Pt(
        DEFAULT_FONT_SIZE
    )

    return y1


def create_word_from_pdf(
    pdf_document,
    translations
):
    pages, elements = (
        extract_pdf_elements(
            pdf_document
        )
    )

    document = Document()

    # python-docx otomatik olarak
    # ilk section ve boş paragraph oluşturur.
    first_paragraph = (
        document.paragraphs[0]
        if document.paragraphs
        else None
    )

    if first_paragraph is not None:
        first_paragraph.text = ""

        first_paragraph.paragraph_format.space_before = Pt(
            0
        )

        first_paragraph.paragraph_format.space_after = Pt(
            0
        )

    for page_index, page_data in enumerate(
        pages
    ):

        page_width = float(
            page_data[
                "width"
            ]
        )

        page_height = float(
            page_data[
                "height"
            ]
        )

        if page_index == 0:
            section = (
                document.sections[0]
            )
        else:
            section = (
                document.add_section(
                    WD_SECTION.NEW_PAGE
                )
            )

        configure_section(
            section,
            page_width,
            page_height
        )

        page_elements = (
            page_data[
                "elements"
            ]
        )

        # PDF üzerindeki konuma göre
        # üstten alta, soldan sağa sırala.
        page_elements = sorted(
            page_elements,
            key=lambda item: (
                float(
                    item[
                        "bbox"
                    ][
                        "y0"
                    ]
                ),
                float(
                    item[
                        "bbox"
                    ][
                        "x0"
                    ]
                )
            )
        )

        previous_bottom = None

        for element in page_elements:

            element_id = (
                element[
                    "id"
                ]
            )

            translated_text = (
                translations.get(
                    element_id,
                    element[
                        "text"
                    ]
                )
            )

            previous_bottom = (
                add_text_block(
                    document,
                    element,
                    translated_text,
                    page_width,
                    previous_bottom
                )
            )

    output = BytesIO()

    document.save(
        output
    )

    output.seek(
        0
    )

    return output


# ============================================================
# EXTRACT PDF
# ============================================================

@app.route(
    "/extract-pdf",
    methods=["POST"]
)
def extract_pdf():

    if "file" not in request.files:
        return jsonify({
            "error":
                "No PDF file received"
        }), 400

    uploaded_file = (
        request.files[
            "file"
        ]
    )

    try:
        pdf_bytes = (
            uploaded_file.read()
        )

        document = fitz.open(
            stream=pdf_bytes,
            filetype="pdf"
        )

    except Exception as e:
        return jsonify({
            "error":
                "Could not read PDF file",

            "details":
                str(e)
        }), 400

    try:
        pages, elements = (
            extract_pdf_elements(
                document
            )
        )

        chunks = [
            elements[
                i:i + CHUNK_SIZE
            ]
            for i in range(
                0,
                len(elements),
                CHUNK_SIZE
            )
        ]

        response = {
            "filename":
                uploaded_file.filename,

            "page_count":
                len(pages),

            "element_count":
                len(elements),

            "chunk_size":
                CHUNK_SIZE,

            "chunk_count":
                len(chunks),

            "pages":
                pages,

            "elements":
                elements,

            "chunks":
                chunks
        }

        return jsonify(
            response
        )

    finally:
        document.close()


# ============================================================
# CREATE DOCX
# ============================================================

@app.route(
    "/create-docx",
    methods=["POST"]
)
def create_docx():

    # --------------------------------------------------------
    # PDF CHECK
    # --------------------------------------------------------

    if "file" not in request.files:
        return jsonify({
            "error":
                "No PDF file received"
        }), 400

    uploaded_file = (
        request.files[
            "file"
        ]
    )

    # --------------------------------------------------------
    # TRANSLATIONS CHECK
    # --------------------------------------------------------

    translations_raw = (
        request.form.get(
            "translations"
        )
    )

    if not translations_raw:
        return jsonify({
            "error":
                "No translations JSON received"
        }), 400

    # --------------------------------------------------------
    # PARSE TRANSLATIONS
    # --------------------------------------------------------

    try:
        translations = (
            parse_translations(
                translations_raw
            )
        )

    except Exception as e:
        return jsonify({
            "error":
                "Could not parse translations JSON",

            "details":
                str(e)
        }), 400

    if not translations:
        return jsonify({
            "error":
                "Translations list is empty"
        }), 400

    # --------------------------------------------------------
    # OPEN ORIGINAL PDF
    # --------------------------------------------------------

    try:
        pdf_bytes = (
            uploaded_file.read()
        )

        pdf_document = fitz.open(
            stream=pdf_bytes,
            filetype="pdf"
        )

    except Exception as e:
        return jsonify({
            "error":
                "Could not read original PDF",

            "details":
                str(e)
        }), 400

    # --------------------------------------------------------
    # CREATE WORD
    # --------------------------------------------------------

    try:
        word_file = (
            create_word_from_pdf(
                pdf_document,
                translations
            )
        )

    except Exception as e:
        return jsonify({
            "error":
                "Could not create DOCX",

            "details":
                str(e)
        }), 500

    finally:
        pdf_document.close()

    # --------------------------------------------------------
    # OUTPUT FILENAME
    # --------------------------------------------------------

    original_name = (
        uploaded_file.filename
        or "translated.pdf"
    )

    if original_name.lower().endswith(
        ".pdf"
    ):
        original_name = (
            original_name[:-4]
        )

    output_filename = (
        f"{original_name}_translated.docx"
    )

    # --------------------------------------------------------
    # RETURN DOCX
    # --------------------------------------------------------

    return send_file(
        word_file,
        as_attachment=True,
        download_name=output_filename,
        mimetype=(
            "application/"
            "vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        )
    )


# ============================================================
# LOCAL
# ============================================================

if __name__ == "__main__":
    app.run()